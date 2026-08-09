"""
NICARA — Export and deliverable-approval tests.

Run: python manage.py test projects.test_exports
"""
import io
from decimal import Decimal

import openpyxl
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import User
from projects.models import (
    BookingForm, Estimate, EstimateItem, PaymentMilestone, Project, ProjectDeliverable,
)

PASSWORD = 'Str0ng!Passw0rd'


class ExportTestCase(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email='admin@nicara.design', password=PASSWORD,
            first_name='NICARA', last_name='Admin', role=User.Role.ADMIN,
        )
        self.client.force_authenticate(self.user)
        self.project = Project.objects.create(
            name='Sharma Residence', client_name='Ms. Anita Sharma',
            client_phone='+91 98765 43210', client_email='anita@example.com',
            developer='Prestige Lakeside', unit_no='B-1204',
            city='Mumbai', area='1,850 sqft',
        )
        self.estimate = Estimate.objects.create(
            project=self.project, type=Estimate.EstimateType.INITIAL, version=1,
        )
        EstimateItem.objects.create(
            estimate=self.estimate, area='Master Bedroom', item='Wardrobe',
            description='16mm BWP ply', length="8'0\"", breadth="2'0\"", height="8'0\"",
            qty=1, unit='unit', rate=Decimal('145000'), gst_pct=18,
        )
        EstimateItem.objects.create(
            estimate=self.estimate, area='Kitchen', item='Base Cabinets',
            qty=1, unit='unit', rate=Decimal('225000'), gst_pct=18,
        )


class EstimatePDFTests(ExportTestCase):
    def url(self):
        return reverse('estimate-pdf', args=[self.project.pk, self.estimate.pk])

    def test_returns_a_pdf(self):
        res = self.client.get(self.url())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res['Content-Type'], 'application/pdf')
        self.assertTrue(res.content.startswith(b'%PDF-'))
        self.assertGreater(len(res.content), 2000)

    def test_filename_is_attached(self):
        res = self.client.get(self.url())
        self.assertIn('attachment;', res['Content-Disposition'])
        self.assertIn('Sharma_Residence', res['Content-Disposition'])
        self.assertIn('.pdf', res['Content-Disposition'])

    def test_requires_authentication(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(self.url()).status_code,
                         status.HTTP_401_UNAUTHORIZED)

    def test_renders_with_a_payment_schedule(self):
        PaymentMilestone.objects.create(
            project=self.project, milestone='Booking Advance',
            amount=Decimal('50000'), due_date='2026-06-15',
        )
        res = self.client.get(self.url())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertTrue(res.content.startswith(b'%PDF-'))

    def test_renders_an_empty_estimate(self):
        empty = Estimate.objects.create(
            project=self.project, type=Estimate.EstimateType.FINAL, version=1,
        )
        res = self.client.get(reverse('estimate-pdf', args=[self.project.pk, empty.pk]))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertTrue(res.content.startswith(b'%PDF-'))

    def test_unknown_estimate_is_404(self):
        res = self.client.get(reverse('estimate-pdf', args=[self.project.pk, 99999]))
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)


class EstimateExcelTests(ExportTestCase):
    def url(self):
        return reverse('estimate-excel', args=[self.project.pk, self.estimate.pk])

    def test_returns_a_workbook(self):
        res = self.client.get(self.url())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', res['Content-Type'])
        self.assertTrue(res.content.startswith(b'PK'))

    def test_workbook_opens_and_carries_the_numbers(self):
        res = self.client.get(self.url())
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb['Estimate']

        text = [
            str(ws.cell(row=r, column=c).value)
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=r, column=c).value is not None
        ]
        self.assertIn('Ms. Anita Sharma', text)
        self.assertIn('Sharma Residence', text)
        self.assertIn('Wardrobe', text)
        self.assertIn('GRAND TOTAL', text)

        # Grand total = (145000 + 225000) × 1.18 must land as a real number so
        # the workbook can be summed, not just read. openpyxl hands whole
        # values back as int, so accept either.
        numbers = [
            float(ws.cell(row=r, column=c).value)
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(row=r, column=c).value, (int, float))
            and not isinstance(ws.cell(row=r, column=c).value, bool)
        ]
        self.assertIn(float(self.estimate.grand_total), numbers)

    def test_payment_sheet_appears_only_when_there_are_milestones(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get(self.url()).content))
        self.assertEqual(wb.sheetnames, ['Estimate'])

        PaymentMilestone.objects.create(
            project=self.project, milestone='Booking Advance',
            amount=Decimal('50000'), due_date='2026-06-15',
        )
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get(self.url()).content))
        self.assertIn('Payment Schedule', wb.sheetnames)


class BookingPDFTests(ExportTestCase):
    def test_returns_a_pdf(self):
        booking = BookingForm.objects.create(
            project=self.project, total_value=Decimal('436600'),
            advance_amount=Decimal('50000'),
        )
        res = self.client.get(reverse('booking-form-pdf', args=[self.project.pk]))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertTrue(res.content.startswith(b'%PDF-'))
        self.assertIn(booking.booking_number, res['Content-Disposition'])

    def test_404_when_no_booking_exists(self):
        res = self.client.get(reverse('booking-form-pdf', args=[self.project.pk]))
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)


class DeliverableVersioningTests(ExportTestCase):
    """FL and Mood Board version control plus the approval workflow."""

    def create_version(self, deliverable_type='furniture_layout', **extra):
        return self.client.post(
            reverse('deliverable-list', args=[self.project.pk]),
            {'type': deliverable_type, 'remarks': 'First cut', **extra},
            format='json',
        )

    def test_version_numbers_increment(self):
        first = self.create_version()
        second = self.create_version()
        self.assertEqual(first.data['version_no'], 1)
        self.assertEqual(second.data['version_no'], 2)
        self.assertEqual(second.data['version'], 'Ver 2')

    def test_only_the_newest_version_is_current(self):
        self.create_version()
        self.create_version()
        current = ProjectDeliverable.objects.filter(
            project=self.project, type='furniture_layout', is_current=True
        )
        self.assertEqual(current.count(), 1)
        self.assertEqual(current.first().version_no, 2)

    def test_new_version_records_what_it_supersedes(self):
        first = self.create_version()
        second = self.create_version()
        superseded = ProjectDeliverable.objects.get(pk=second.data['id']).supersedes
        self.assertEqual(superseded.pk, first.data['id'])

    def test_types_version_independently(self):
        self.create_version('furniture_layout')
        mood = self.create_version('mood_board')
        self.assertEqual(mood.data['version_no'], 1)

    def test_starts_as_draft(self):
        res = self.create_version()
        self.assertEqual(res.data['status'], 'draft')

    def test_submit_for_approval(self):
        created = self.create_version()
        res = self.client.post(
            reverse('deliverable-submit', args=[self.project.pk, created.data['id']])
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['status'], 'pending')
        self.assertIsNotNone(res.data['submitted_at'])
        self.assertEqual(res.data['submitted_by_name'], 'NICARA Admin')

    def test_approve_records_the_reviewer(self):
        created = self.create_version()
        res = self.client.post(
            reverse('deliverable-approve', args=[self.project.pk, created.data['id']]),
            {'remarks': 'Looks good'}, format='json',
        )
        self.assertEqual(res.data['status'], 'approved')
        self.assertEqual(res.data['review_remarks'], 'Looks good')
        self.assertEqual(res.data['reviewed_by_name'], 'NICARA Admin')
        self.assertIsNotNone(res.data['reviewed_at'])

    def test_approving_an_older_version_makes_it_current(self):
        first = self.create_version()
        self.create_version()  # v2 becomes current on upload
        self.client.post(
            reverse('deliverable-approve', args=[self.project.pk, first.data['id']]),
            {'remarks': 'Client preferred v1'}, format='json',
        )
        current = ProjectDeliverable.objects.filter(
            project=self.project, type='furniture_layout', is_current=True
        )
        self.assertEqual(current.count(), 1)
        self.assertEqual(current.first().pk, first.data['id'])

    def test_request_revision_needs_remarks(self):
        created = self.create_version()
        res = self.client.post(
            reverse('deliverable-revision', args=[self.project.pk, created.data['id']]),
            {'remarks': '   '}, format='json',
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_request_revision_stores_the_feedback(self):
        created = self.create_version()
        res = self.client.post(
            reverse('deliverable-revision', args=[self.project.pk, created.data['id']]),
            {'remarks': 'Kitchen island position needs to change'}, format='json',
        )
        self.assertEqual(res.data['status'], 'revision')
        self.assertIn('Kitchen island', res.data['review_remarks'])

    def test_cannot_resubmit_an_approved_version(self):
        created = self.create_version()
        self.client.post(
            reverse('deliverable-approve', args=[self.project.pk, created.data['id']]),
            {}, format='json',
        )
        res = self.client.post(
            reverse('deliverable-submit', args=[self.project.pk, created.data['id']])
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_filter_by_current(self):
        self.create_version()
        self.create_version()
        res = self.client.get(
            reverse('deliverable-list', args=[self.project.pk]), {'is_current': 'true'}
        )
        self.assertEqual(res.data['count'], 1)
